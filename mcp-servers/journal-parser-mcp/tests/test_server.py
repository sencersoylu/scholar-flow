"""Tests for Journal Parser MCP Server."""

import tempfile
from pathlib import Path

import pytest
import pytest_asyncio

from docx import Document
from docx.shared import Pt, Inches

from journal_parser_mcp.server import mcp, parse_docx_template, parse_latex_template, parse_guidelines_text
from journal_parser_mcp.parsers import parse_docx, parse_latex, parse_guidelines_text as parse_text
from journal_parser_mcp.models import JournalProfile, FormattingRules, WordLimits, FigureRequirements


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Create a sample .docx template for testing."""
    doc = Document()

    # Set page margins
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

    # Add headings and body text
    doc.add_heading("Introduction", level=1)
    para = doc.add_paragraph("This is body text in the introduction section.")
    run = para.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    doc.add_heading("Methods", level=1)
    doc.add_paragraph("Description of methods used.")

    doc.add_heading("Data Collection", level=2)
    doc.add_paragraph("Details on data collection.")

    doc.add_heading("Results", level=1)
    doc.add_paragraph("Key findings are presented here.")

    doc.add_heading("Discussion", level=1)
    doc.add_paragraph("Interpretation of results.")

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("Summary and future directions.")

    path = tmp_path / "sample_template.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def sample_latex(tmp_path: Path) -> Path:
    """Create a sample LaTeX file for testing."""
    content = r"""\documentclass[12pt,twocolumn]{article}

\usepackage[a4paper,top=1in,bottom=1in,left=1.25in,right=1.25in]{geometry}
\usepackage{times}
\usepackage{graphicx}
\usepackage{amsmath,amssymb}
\usepackage[numbers]{natbib}

\linespread{1.5}

\bibliographystyle{unsrt}

\begin{document}

\title{Sample Article}
\author{Test Author}
\maketitle

\section{Introduction}
This is the introduction.

\section{Methods}
\subsection{Data Collection}
Details of data collection.

\section{Results}
Key findings.

\section{Discussion}
Interpretation.

\section{Conclusion}
Summary.

\bibliography{references}

\end{document}
"""
    path = tmp_path / "sample.tex"
    path.write_text(content)
    return path


@pytest.fixture
def sample_guidelines_text() -> str:
    """Sample author guidelines text."""
    return """
    Author Guidelines for the Journal of Example Research

    Manuscripts should not exceed 5000 words. The abstract must be limited to 250 words.
    Title should not exceed 15 words.

    Please provide 3 to 6 keywords.

    Required Sections: Abstract, Introduction, Methods, Results, Discussion,
    Conclusion, Acknowledgments, References.

    References should be numbered in order of appearance (Vancouver style).

    Figures should be submitted in TIFF or EPS format at a minimum of 300 DPI.
    Maximum of 8 figures per manuscript. Each figure file must not exceed 10 MB.

    Authors must include a Data Availability statement and Conflict of Interest declaration.
    """


# ---------------------------------------------------------------------------
# Server registration tests
# ---------------------------------------------------------------------------

def test_server_name():
    assert mcp.name == "journal-parser-mcp"


def test_tools_registered():
    assert len(mcp._tool_manager._tools) >= 3


# ---------------------------------------------------------------------------
# DOCX parsing tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_parse_docx_template(sample_docx: Path):
    result = await parse_docx_template(str(sample_docx))
    assert "# Journal Profile:" in result
    assert "Introduction" in result
    assert "Methods" in result
    assert "Results" in result
    assert "Discussion" in result
    assert "Conclusion" in result


@pytest.mark.asyncio
async def test_parse_docx_extracts_formatting(sample_docx: Path):
    profile = parse_docx(str(sample_docx))
    assert profile.formatting.margins_inches.get("top") == 1.0
    assert profile.formatting.margins_inches.get("left") == 1.25
    assert len(profile.sections) == 6  # Intro, Methods, Data Collection, Results, Discussion, Conclusion


@pytest.mark.asyncio
async def test_parse_docx_font_info(sample_docx: Path):
    profile = parse_docx(str(sample_docx))
    assert profile.formatting.font_family == "Times New Roman"
    assert profile.formatting.font_size_pt == 12.0


@pytest.mark.asyncio
async def test_parse_docx_file_not_found():
    result = await parse_docx_template("/nonexistent/path/file.docx")
    assert "**Error:**" in result
    assert "not found" in result.lower()


@pytest.mark.asyncio
async def test_parse_docx_invalid_format(tmp_path: Path):
    bad_file = tmp_path / "readme.txt"
    bad_file.write_text("not a docx")
    result = await parse_docx_template(str(bad_file))
    assert "**Error:**" in result


# ---------------------------------------------------------------------------
# LaTeX parsing tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_parse_latex_template(sample_latex: Path):
    result = await parse_latex_template(str(sample_latex))
    assert "# Journal Profile:" in result
    assert "Introduction" in result
    assert "Methods" in result
    assert "article" in result  # document class


@pytest.mark.asyncio
async def test_parse_latex_extracts_details(sample_latex: Path):
    profile = parse_latex(str(sample_latex))
    assert profile.document_class == "article"
    assert profile.formatting.font_size_pt == 12.0
    assert profile.formatting.columns == 2
    assert profile.formatting.page_size == "A4"
    assert profile.citation_style == "numbered (natbib)"
    assert profile.bibliography_style == "unsrt"
    assert profile.formatting.line_spacing == 1.5
    assert "graphicx" in profile.packages
    assert "amsmath" in profile.packages


@pytest.mark.asyncio
async def test_parse_latex_margins(sample_latex: Path):
    profile = parse_latex(str(sample_latex))
    assert profile.formatting.margins_inches.get("top") == 1.0
    assert profile.formatting.margins_inches.get("left") == 1.25


@pytest.mark.asyncio
async def test_parse_latex_file_not_found():
    result = await parse_latex_template("/nonexistent/file.tex")
    assert "**Error:**" in result


@pytest.mark.asyncio
async def test_parse_latex_invalid_format(tmp_path: Path):
    bad = tmp_path / "data.csv"
    bad.write_text("a,b,c")
    result = await parse_latex_template(str(bad))
    assert "**Error:**" in result


# ---------------------------------------------------------------------------
# Guidelines text parsing tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_parse_guidelines_text(sample_guidelines_text: str):
    result = await parse_guidelines_text(sample_guidelines_text)
    assert "# Journal Profile:" in result
    assert "5000" in result  # manuscript word limit
    assert "250" in result   # abstract word limit


@pytest.mark.asyncio
async def test_parse_guidelines_extracts_details(sample_guidelines_text: str):
    profile = parse_text(sample_guidelines_text)
    assert profile.word_limits.abstract == 250
    assert profile.word_limits.manuscript == 5000
    assert profile.word_limits.title == 15
    assert profile.word_limits.keywords_count == 6
    assert profile.figure_requirements.min_resolution_dpi == 300
    assert profile.figure_requirements.max_count == 8
    assert profile.figure_requirements.max_file_size_mb == 10.0
    assert "TIFF" in profile.figure_requirements.formats
    assert "EPS" in profile.figure_requirements.formats
    assert "Vancouver" in (profile.citation_style or "")
    assert "Introduction" in profile.sections
    assert "Methods" in profile.sections
    assert "Data Availability" in profile.sections


@pytest.mark.asyncio
async def test_parse_guidelines_empty_text():
    result = await parse_guidelines_text("")
    assert "**Error:**" in result


@pytest.mark.asyncio
async def test_parse_guidelines_minimal_text():
    profile = parse_text("Please submit your manuscript with an abstract of 300 words.")
    assert profile.word_limits.abstract == 300


# ---------------------------------------------------------------------------
# Model tests
# ---------------------------------------------------------------------------

def test_journal_profile_to_markdown():
    profile = JournalProfile(
        name="Test Journal",
        formatting=FormattingRules(
            font_family="Arial",
            font_size_pt=11.0,
            columns=1,
        ),
        sections=["Introduction", "Methods", "Results"],
        citation_style="APA",
        word_limits=WordLimits(abstract=200, manuscript=4000),
        figure_requirements=FigureRequirements(formats=["PNG", "TIFF"], min_resolution_dpi=300),
    )
    md = profile.to_markdown()
    assert "# Journal Profile: Test Journal" in md
    assert "Arial" in md
    assert "11.0 pt" in md
    assert "1. Introduction" in md
    assert "200 words" in md
    assert "4000 words" in md
    assert "300 DPI" in md
    assert "APA" in md


def test_journal_profile_empty():
    profile = JournalProfile(name="Empty")
    md = profile.to_markdown()
    assert "# Journal Profile: Empty" in md
    # Should still render without errors
    assert "## Formatting Rules" in md

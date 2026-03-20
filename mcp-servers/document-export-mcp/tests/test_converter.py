"""Tests for the converter module."""

import tempfile
from pathlib import Path

import pytest
from docx import Document

from document_export_mcp.converter import (
    _parse_markdown,
    _parse_table_lines,
    markdown_to_docx,
    markdown_to_latex,
)

SAMPLE_MD = """\
# My Research Paper

## Introduction

This is the **introduction** with *italic* text and ***bold italic***.

## Methods

We used a novel approach.

### Data Collection

Data was collected from multiple sources \\cite{smith2024}.

```python
import pandas as pd
df = pd.read_csv("data.csv")
```

## Results

| Metric | Value | P-value |
|--------|-------|---------|
| Accuracy | 0.95 | 0.001 |
| Recall | 0.92 | 0.003 |

## Conclusion

The results are significant.
"""


@pytest.fixture
def sample_md(tmp_path: Path) -> Path:
    md_file = tmp_path / "manuscript.md"
    md_file.write_text(SAMPLE_MD)
    return md_file


class TestParseMarkdown:
    def test_parse_headings(self):
        blocks = _parse_markdown("# Title\n\n## Section\n\n### Subsection")
        headings = [b for b in blocks if b["type"] == "heading"]
        assert len(headings) == 3
        assert headings[0]["level"] == 1
        assert headings[0]["text"] == "Title"
        assert headings[1]["level"] == 2
        assert headings[2]["level"] == 3

    def test_parse_paragraphs(self):
        blocks = _parse_markdown("# Title\n\nThis is a paragraph.\n\nAnother paragraph.")
        paragraphs = [b for b in blocks if b["type"] == "paragraph"]
        assert len(paragraphs) == 2
        assert paragraphs[0]["text"] == "This is a paragraph."

    def test_parse_code_block(self):
        md = "```python\nprint('hello')\n```"
        blocks = _parse_markdown(md)
        code_blocks = [b for b in blocks if b["type"] == "code_block"]
        assert len(code_blocks) == 1
        assert code_blocks[0]["lang"] == "python"
        assert "print('hello')" in code_blocks[0]["content"]

    def test_parse_table(self):
        md = "| A | B |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |"
        blocks = _parse_markdown(md)
        tables = [b for b in blocks if b["type"] == "table"]
        assert len(tables) == 1
        headers, rows = _parse_table_lines(tables[0]["lines"])
        assert headers == ["A", "B"]
        assert len(rows) == 2

    def test_parse_full_sample(self):
        blocks = _parse_markdown(SAMPLE_MD)
        types = [b["type"] for b in blocks]
        assert "heading" in types
        assert "paragraph" in types
        assert "code_block" in types
        assert "table" in types


class TestMarkdownToDocx:
    def test_basic_conversion(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.docx"
        result = markdown_to_docx(str(sample_md), str(output))
        assert result == str(output)
        assert output.exists()
        assert output.stat().st_size > 0

    def test_docx_readable(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.docx"
        markdown_to_docx(str(sample_md), str(output))
        doc = Document(str(output))
        # Should have paragraphs
        assert len(doc.paragraphs) > 0
        # Check that at least one table exists
        assert len(doc.tables) >= 1

    def test_docx_table_content(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.docx"
        markdown_to_docx(str(sample_md), str(output))
        doc = Document(str(output))
        table = doc.tables[0]
        # Header row
        assert table.rows[0].cells[0].text == "Metric"
        # Data row
        assert table.rows[1].cells[1].text == "0.95"

    def test_file_not_found(self, tmp_path: Path):
        with pytest.raises(FileNotFoundError):
            markdown_to_docx(str(tmp_path / "nonexistent.md"), str(tmp_path / "out.docx"))

    def test_with_template(self, sample_md: Path, tmp_path: Path):
        # Create a minimal template
        template = tmp_path / "template.docx"
        doc = Document()
        doc.save(str(template))

        output = tmp_path / "output.docx"
        result = markdown_to_docx(str(sample_md), str(output), template_path=str(template))
        assert output.exists()

    def test_creates_output_dirs(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "nested" / "dirs" / "output.docx"
        markdown_to_docx(str(sample_md), str(output))
        assert output.exists()


class TestMarkdownToLatex:
    def test_basic_conversion(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.tex"
        result = markdown_to_latex(str(sample_md), str(output))
        assert result == str(output)
        assert output.exists()

    def test_latex_structure(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.tex"
        markdown_to_latex(str(sample_md), str(output))
        content = output.read_text()
        assert "\\documentclass{article}" in content
        assert "\\begin{document}" in content
        assert "\\end{document}" in content
        assert "\\section{" in content
        assert "\\subsection{" in content
        assert "\\subsubsection{" in content

    def test_latex_inline_formatting(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.tex"
        markdown_to_latex(str(sample_md), str(output))
        content = output.read_text()
        assert "\\textbf{" in content
        assert "\\textit{" in content

    def test_latex_code_block(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.tex"
        markdown_to_latex(str(sample_md), str(output))
        content = output.read_text()
        assert "\\begin{verbatim}" in content
        assert "\\end{verbatim}" in content

    def test_latex_table(self, sample_md: Path, tmp_path: Path):
        output = tmp_path / "output.tex"
        markdown_to_latex(str(sample_md), str(output))
        content = output.read_text()
        assert "\\begin{tabular}" in content
        assert "\\end{tabular}" in content
        assert "\\toprule" in content

    def test_latex_citations(self, tmp_path: Path):
        md_file = tmp_path / "cite.md"
        md_file.write_text("# Title\n\nAs shown by \\cite{smith2024}.")
        output = tmp_path / "output.tex"
        markdown_to_latex(str(md_file), str(output))
        content = output.read_text()
        assert "\\cite{smith2024}" in content

    def test_file_not_found(self, tmp_path: Path):
        with pytest.raises(FileNotFoundError):
            markdown_to_latex(str(tmp_path / "nonexistent.md"), str(tmp_path / "out.tex"))

    def test_with_template(self, sample_md: Path, tmp_path: Path):
        template = tmp_path / "template.tex"
        template.write_text("\\documentclass{report}\n\\usepackage{amsmath}\n")
        output = tmp_path / "output.tex"
        markdown_to_latex(str(sample_md), str(output), template_path=str(template))
        content = output.read_text()
        assert "\\documentclass{report}" in content
        assert "\\usepackage{amsmath}" in content
